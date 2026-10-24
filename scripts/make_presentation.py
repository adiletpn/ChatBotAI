"""Генерирует презентацию проекта docs/DocMind_AI_Presentation.pptx (16:9, с заметками докладчика).

Запуск:  .venv/bin/python scripts/make_presentation.py
Поля «Команда / Колледж / Руководитель» на титульном слайде заполните своими данными.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"

TEAL = RGBColor(0xC2, 0x60, 0x3E)  # фирменный терракотовый
TEAL_SOFT = RGBColor(0xF6, 0xE8, 0xE1)
DARK = RGBColor(0x1F, 0x1E, 0x1D)
MUTED = RGBColor(0x75, 0x73, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE6, 0xE3, 0xD8)
FONT = "Arial"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, value, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = value if isinstance(value, list) else [value]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(size), bold, color, FONT
    return box


def bullets(slide, x, y, w, h, items, size=18, gap=10):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        head, _, tail = item.partition(" — ")
        r = p.add_run()
        r.text = "▪  " + head
        r.font.size, r.font.color.rgb, r.font.name = Pt(size), DARK, FONT
        r.font.bold = bool(tail)
        if tail:
            r2 = p.add_run()
            r2.text = " — " + tail
            r2.font.size, r2.font.color.rgb, r2.font.name = Pt(size), MUTED, FONT
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def slide_base(title, number=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    number = len(prs.slides)
    rect(s, 0, 0, Inches(0.18), H, fill=TEAL, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, Inches(0.6), Inches(0.35), Inches(11.5), Inches(0.9), title, size=30, bold=True)
    text(s, W - Inches(1.6), H - Inches(0.55), Inches(1.2), Inches(0.4), f"DocMind AI · {number}", size=10,
         color=MUTED, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def card(slide, x, y, w, h, icon, title, body, title_size=17, body_size=13):
    rect(slide, x, y, w, h)
    text(slide, x + Inches(0.2), y + Inches(0.15), w - Inches(0.4), Inches(0.5), icon, size=24)
    text(slide, x + Inches(0.2), y + Inches(0.7), w - Inches(0.4), Inches(0.5), title, size=title_size, bold=True)
    text(slide, x + Inches(0.2), y + Inches(1.15), w - Inches(0.4), h - Inches(1.2), body, size=body_size, color=MUTED)


def picture(slide, path, x, y, w=None, h=None):
    pic = slide.shapes.add_picture(str(path), x, y, w, h)
    pic.line.color.rgb = BORDER
    return pic


# 1. Титул ---------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, fill=TEAL, line=None, shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(0.8), Inches(0.7), Inches(8), Inches(0.5), "INNOVATION AI HACKATHON – 2026  ·  AIChatBots", size=16,
     bold=True, color=RGBColor(0xF6, 0xE8, 0xE1))
text(s, Inches(0.8), Inches(1.7), Inches(7.5), Inches(1.4), "DocMind AI", size=66, bold=True, color=WHITE)
text(s, Inches(0.8), Inches(3.0), Inches(6.6), Inches(1.6),
     "Құжаттармен сөйлесетін зияткер чат-бот\nИнтеллектуальный чат-бот, который отвечает на вопросы по документам",
     size=20, color=WHITE)
text(s, Inches(0.8), Inches(5.3), Inches(6.6), Inches(1.4),
     ["Команда: ____________", "Колледж: ____________", "Руководитель: ____________"], size=16, color=WHITE)
picture(s, DOCS / "ui_mobile_dark.png", Inches(9.3), Inches(0.6), h=Inches(6.3))
s.notes_slide.notes_text_frame.text = (
    "Здравствуйте! Мы команда ___ из ___. Представляем DocMind AI — чат-бот, который читает ваши документы "
    "и отвечает на вопросы по ним на казахском, русском и английском, всегда указывая источник.")

# 2. Проблема --------------------------------------------------------------
s = slide_base("Проблема", 2, "Расскажите о проблеме на живом примере: студент ищет в регламенте дедлайн, "
               "рабочий — норму выдачи СИЗ в инструкции. Документы длинные, на разных языках, часть — сканы.")
problems = [
    ("⏳", "Долгий поиск", "Регламенты, инструкции и приказы занимают десятки страниц — нужный пункт ищут вручную."),
    ("🌐", "Разные языки", "Документы на казахском, вопросы — на русском или английском, и наоборот."),
    ("🖨️", "Сканы и фото", "Многие документы существуют только как сканы — по ним не работает обычный поиск Ctrl+F."),
    ("❓", "Нет доверия к ИИ", "Обычные чат-боты «придумывают» ответы и не показывают, откуда взята информация."),
]
for i, (icon, title, body) in enumerate(problems):
    card(s, Inches(0.6) + i * Inches(3.08), Inches(1.6), Inches(2.9), Inches(3.3), icon, title, body)
text(s, Inches(0.6), Inches(5.4), Inches(12), Inches(1),
     "Особенно остро это на производстве: инструкции по охране труда, технологические регламенты и нормативы "
     "должны находиться быстро и без ошибок.", size=17, color=MUTED)

# 3. Решение -------------------------------------------------------------
s = slide_base("Решение — DocMind AI", 3, "Покажите три шага. Главное отличие — ответ всегда с источником: "
               "документ, страница и цитата, которую можно открыть.")
steps = [("1", "Загрузите документ", "PDF, DOCX, PPTX, XLSX, TXT, HTML, фото или скан — через сайт или Telegram"),
         ("2", "Задайте вопрос", "На казахском, русском или английском — текстом или голосом"),
         ("3", "Получите ответ с источником", "Точный ответ по документу + ссылки [1] [2]: файл, страница, цитата")]
for i, (n, title, body) in enumerate(steps):
    y = Inches(1.6) + i * Inches(1.55)
    circle = rect(s, Inches(0.7), y, Inches(0.8), Inches(0.8), fill=TEAL, line=None, shape=MSO_SHAPE.OVAL)
    circle.text_frame.text = n
    p = circle.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.size, p.runs[0].font.bold, p.runs[0].font.color.rgb = Pt(24), True, WHITE
    text(s, Inches(1.75), y - Inches(0.05), Inches(4.6), Inches(0.5), title, size=20, bold=True)
    text(s, Inches(1.75), y + Inches(0.45), Inches(4.6), Inches(0.9), body, size=14, color=MUTED)
picture(s, DOCS / "ui_chat.png", Inches(6.6), Inches(1.5), w=Inches(6.2))

# 3a. Доказательство на странице -------------------------------------------
s = slide_base("Не просто чат: ответ с доказательством на странице", notes=
               "Главная фишка: нажимаем на источник — открывается настоящая страница документа, а строки, "
               "на которых основан ответ, подсвечены маркером. Ответ можно проверить за секунду.")
picture(s, DOCS / "ui_evidence.png", Inches(0.6), Inches(1.4), w=Inches(8.2))
bullets(s, Inches(9.1), Inches(1.6), Inches(3.9), Inches(5), [
    "📍 Дәлел — клик по источнику открывает страницу документа",
    "Жёлтый маркер — подсвечены именно те строки, на которых основан ответ",
    "Индикатор доверия — «подтверждено N источниками» под каждым ответом",
    "Нет доказательства — бот честно предупреждает",
], size=15)

# 3b. Режимы ---------------------------------------------------------------
s = slide_base("Три режима работы с документом", notes=
               "Обычные чат-боты умеют только отвечать на вопросы. DocMind сам анализирует документ "
               "и проверяет, как человек его усвоил — это важно для инструктажей по охране труда.")
modes = [("💬 Вопрос-ответ", "диалог на 3 языках с источниками"),
         ("🧭 Карта документа", "сроки, цифры, чек-лист требований, термины — без единого вопроса"),
         ("🎓 Проверка знаний", "тест по документу с объяснениями и баллами")]
for i, (t, d) in enumerate(modes):
    x = Inches(0.6) + i * Inches(4.1)
    rect(s, x, Inches(1.4), Inches(3.9), Inches(1.05), fill=TEAL_SOFT if i else WHITE)
    text(s, x + Inches(0.15), Inches(1.45), Inches(3.6), Inches(0.45), t, size=17, bold=True)
    text(s, x + Inches(0.15), Inches(1.9), Inches(3.6), Inches(0.5), d, size=12.5, color=MUTED)
picture(s, DOCS / "ui_map.png", Inches(0.6), Inches(2.75), w=Inches(6.0))
picture(s, DOCS / "ui_quiz.png", Inches(6.85), Inches(2.75), w=Inches(6.0))

# 4. Как работает ---------------------------------------------------------
s = slide_base("Как это работает: технология RAG", 4,
               "RAG — Retrieval-Augmented Generation. Модель не «вспоминает», а читает найденные фрагменты документа. "
               "Поэтому ответы точные, а каждый факт подтверждён ссылкой.")
pipeline = [("📄", "Парсинг", "PDF, DOCX, PPTX, XLSX…\nсканы → OCR"),
            ("✂️", "Фрагменты", "разбиение по абзацам,\nномер страницы"),
            ("🔎", "Поиск", "BM25 + стемминг kk/ru,\nперевод запроса ×3"),
            ("🧠", "LLM", "ответ только по\nнайденному тексту"),
            ("✅", "Ответ", "на языке вопроса\n+ источники [n]")]
bw, gap = Inches(2.15), Inches(0.33)
for i, (icon, title, body) in enumerate(pipeline):
    x = Inches(0.6) + i * (bw + gap)
    rect(s, x, Inches(1.8), bw, Inches(2.3), fill=TEAL_SOFT if i in (2, 3) else WHITE)
    text(s, x, Inches(1.95), bw, Inches(0.6), icon, size=28, align=PP_ALIGN.CENTER)
    text(s, x, Inches(2.6), bw, Inches(0.5), title, size=18, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, Inches(3.1), bw, Inches(0.9), body, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    if i < len(pipeline) - 1:
        arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + bw + Inches(0.04), Inches(2.8), Inches(0.25), Inches(0.3))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = TEAL
        arrow.line.fill.background()
bullets(s, Inches(0.6), Inches(4.5), Inches(12), Inches(2.5), [
    "Межъязыковой поиск — вопрос на русском находит ответ в казахском документе",
    "Защита от «галлюцинаций» — промпт запрещает выдумывать; нет ответа в документе → бот честно говорит об этом",
    "Офлайн-режим — без интернета бот показывает самые релевантные предложения с источниками",
], size=16)

# 5. Возможности ----------------------------------------------------------
s = slide_base("Возможности", 5, "Подчеркните, что все три бонусных пункта регламента (источник, многоязычие, сканы) реализованы.")
features = [("📚", "10+ форматов", "PDF, DOCX, PPTX, XLSX, CSV, TXT, HTML, DOC/RTF, JPG/PNG"),
            ("📍", "Источник ответа ★", "файл, страница / слайд / лист, цитата, открытие на нужной странице"),
            ("🌐", "3 языка ★", "қазақша · русский · English — вопросы, ответы и интерфейс"),
            ("🖨️", "OCR сканов ★", "Tesseract kaz+rus+eng: сканы PDF, фото, картинки в DOCX"),
            ("💬", "Веб + Telegram", "сайт с адаптивным дизайном и Telegram-бот"),
            ("🧭", "Карта документа", "хронология сроков, ключевые цифры, чек-лист требований, термины"),
            ("🎙️", "Голосовой ввод", "вопрос голосом на казахском, русском, английском"),
            ("🎓", "Проверка знаний", "тест по документу: баллы, объяснения, разбор ошибок")]
for i, (icon, title, body) in enumerate(features):
    col, row = i % 4, i // 4
    card(s, Inches(0.6) + col * Inches(3.08), Inches(1.5) + row * Inches(2.75), Inches(2.9), Inches(2.5), icon, title, body,
         title_size=16, body_size=12.5)
text(s, Inches(0.6), Inches(6.85), Inches(8), Inches(0.4), "★ — дополнительные возможности, отмеченные в регламенте (п. 5.1.2.7)",
     size=11, color=MUTED)

# 6. Соответствие требованиям ---------------------------------------------
s = slide_base("Соответствие требованиям регламента", 6)
rows = [("Требование п. 5.1.2", "Реализация"),
        ("1. Работа с PDF, DOCX и др.", "10+ форматов, таблицы, сканы"),
        ("2. Точные и понятные ответы", "RAG + запрет выдумывать + цитаты"),
        ("3. ИИ и NLP", "LLM, определение языка, стемминг, BM25, перевод запроса"),
        ("4. Telegram / WhatsApp / веб", "Веб-платформа и Telegram-бот"),
        ("5. Рабочая версия", "Запуск одной командой, автотест: 20 проверок"),
        ("6. Презентация и видео", "Эта презентация + видео 5–7 минут"),
        ("7. Источник, многоязычие, сканы", "Все три реализованы")]
table = s.shapes.add_table(len(rows), 2, Inches(0.6), Inches(1.5), Inches(12.1), Inches(5)).table
table.columns[0].width, table.columns[1].width = Inches(5), Inches(7.1)
for r, (a, b) in enumerate(rows):
    for c, value in enumerate((a, b)):
        cell = table.cell(r, c)
        cell.text = value
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.size, para.runs[0].font.name = Pt(16), FONT
        para.runs[0].font.bold = r == 0 or c == 0
        para.runs[0].font.color.rgb = WHITE if r == 0 else DARK
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEAL if r == 0 else (TEAL_SOFT if r % 2 else WHITE)

# 7. Технологии -----------------------------------------------------------
s = slide_base("Технологии", 7, "Отвечая на вопросы жюри: модель подключается через .env — любой OpenAI- или "
               "Anthropic-совместимый API, можно бесплатный.")
stack = [("Backend", "Python 3, FastAPI, потоковые ответы (SSE)"),
         ("ИИ / LLM", "OpenAI- или Anthropic-совместимый API (подключается через .env)"),
         ("Поиск", "BM25 + стемминг для казахского и русского, межъязыковое расширение запроса"),
         ("Документы", "PyMuPDF, python-docx, python-pptx, openpyxl, BeautifulSoup"),
         ("OCR", "Tesseract 5 с моделями kaz, rus, eng"),
         ("Интерфейсы", "HTML/CSS/JS (адаптивный, тёмная тема), Telegram — aiogram 3"),
         ("Безопасность", "Вход по логину/паролю: хэши PBKDF2, HttpOnly-cookie, изоляция документов")]
for i, (k, v) in enumerate(stack):
    y = Inches(1.45) + i * Inches(0.78)
    rect(s, Inches(0.6), y, Inches(2.6), Inches(0.65), fill=TEAL, line=None)
    text(s, Inches(0.6), y, Inches(2.6), Inches(0.65), k, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(3.4), y, Inches(9.3), Inches(0.65), v, size=17, anchor=MSO_ANCHOR.MIDDLE)

# 8. Демонстрация ---------------------------------------------------------
s = slide_base("Демонстрация", 8, "Здесь переключитесь на живую демонстрацию: загрузите регламент хакатона, "
               "спросите по-русски «Когда финал?», откройте источник, затем загрузите скан.")
picture(s, DOCS / "ui_welcome.png", Inches(0.6), Inches(1.5), w=Inches(7.6))
bullets(s, Inches(8.6), Inches(1.6), Inches(4.3), Inches(5), [
    "Регламент хакатона (PDF, қазақша) — «Когда финал?» по-русски → 14 ноября, стр. 1",
    "Скан без текста — OCR распознаёт таблицу баллов",
    "Инструкция по охране труда (DOCX) — «Срок носки каски?» → 24 месяца",
    "Telegram — тот же бот в мессенджере",
], size=15)

# 9. Применение -----------------------------------------------------------
s = slide_base("Где применять", 9)
uses = [("🏭", "Производство", "инструкции по охране труда, техрегламенты, нормы СИЗ — быстрый ответ рабочему прямо в Telegram"),
        ("🎓", "Колледжи и вузы", "положения, правила приёма, расписания, методички — помощник для студентов"),
        ("🏛️", "Госсектор и бизнес", "приказы, договоры, стандарты — поиск по архиву документов на трёх языках"),
        ("⚖️", "Юристы и HR", "договоры, трудовые нормы, внутренние политики — ответы с точной ссылкой на пункт")]
for i, (icon, title, body) in enumerate(uses):
    col, row = i % 2, i // 2
    card(s, Inches(0.6) + col * Inches(6.15), Inches(1.5) + row * Inches(2.65), Inches(5.95), Inches(2.4), icon, title, body,
         body_size=15)

# 10. Уникальность и развитие ---------------------------------------------
s = slide_base("Чем мы отличаемся и что дальше", 10)
text(s, Inches(0.6), Inches(1.4), Inches(6), Inches(0.5), "Преимущества", size=20, bold=True, color=TEAL)
bullets(s, Inches(0.6), Inches(2.0), Inches(5.9), Inches(4.8), [
    "Казахский язык — стемминг и OCR под казахский",
    "Доказательность — подсветка ответа прямо на странице документа",
    "Не только чат — карта документа и проверка знаний",
    "Работает со сканами — бесплатный локальный OCR",
    "Без интернета — офлайн-режим поиска",
    "Свободный выбор модели — любой совместимый API",
], size=16)
text(s, Inches(6.9), Inches(1.4), Inches(6), Inches(0.5), "План развития", size=20, bold=True, color=TEAL)
bullets(s, Inches(6.9), Inches(2.0), Inches(5.9), Inches(4.8), [
    "WhatsApp-бот",
    "Векторный (семантический) поиск",
    "Подключение к системам документооборота",
    "Голосовые ответы на казахском",
    "Роли и общие базы знаний для организаций",
], size=16)

# 11. Спасибо -------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, fill=TEAL, line=None, shape=MSO_SHAPE.RECTANGLE)
text(s, 0, Inches(2.4), W, Inches(1.2), "Рахмет! Спасибо!", size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
text(s, 0, Inches(3.7), W, Inches(1), "DocMind AI — ответы по документам с точным источником", size=22, color=WHITE,
     align=PP_ALIGN.CENTER)
text(s, 0, Inches(5.0), W, Inches(1), "Сайт: ____________     Telegram: @____________", size=18,
     color=RGBColor(0xF6, 0xE8, 0xE1), align=PP_ALIGN.CENTER)

out = DOCS / "DocMind_AI_Presentation.pptx"
prs.save(out)
print("saved", out)
