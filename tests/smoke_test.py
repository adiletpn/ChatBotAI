"""Сквозная проверка без LLM: загрузка демо-документов, поиск, офлайн-ответы, API.

Запуск:  .venv/bin/python tests/smoke_test.py
"""
import asyncio
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="docmind_test_")
os.environ["LLM_API_KEY"] = ""  # проверяем офлайн-режим
os.environ["TELEGRAM_BOT_TOKEN"] = ""
os.environ["SMTP_HOST"] = ""  # коды — в консоль, настоящие письма в тестах не шлём
os.environ["SMTP_USER"] = ""
os.environ["BREVO_API_KEY"] = ""
os.environ["MODELS_DIR"] = str(Path.home() / ".cache" / "mind-ai-models")  # модель эмбеддингов качаем один раз

from fastapi.testclient import TestClient  # noqa: E402

from app import rag  # noqa: E402
from app.config import settings  # noqa: E402
from app.lang import detect_lang  # noqa: E402
from app.main import app  # noqa: E402
from app.ocr import ocr_engine  # noqa: E402
from app.store import store  # noqa: E402

SAMPLES = ROOT / "samples"
failures = 0


def check(name: str, ok: bool, detail: str = "") -> None:
    global failures
    print(("  ✅ " if ok else "  ❌ ") + name + (f" — {detail}" if detail else ""))
    failures += 0 if ok else 1


async def ask(question: str, owner: str) -> dict:
    return await rag.answer(question, owner, owner + ":" + question)


async def main() -> None:
    settings.llm_api_key = ""
    print("Язык:")
    check("kk", detect_lang("Финал қашан өтеді?") == "kk")
    check("kk без особых букв", detect_lang("Командада неше адам болады?") == "kk")
    check("ru", detect_lang("Когда пройдёт финал?") == "ru")
    check("en", detect_lang("When is the final?") == "en")

    print("Загрузка документов:")
    owner = "test"
    pdf = rag.ingest(owner, "Ereje.pdf", (SAMPLES / "Innovation_AI_Hackathon_2026_Ereje.pdf").read_bytes())
    check("PDF", pdf["pages"] == 2 and pdf["chunks"] > 5 and pdf["lang"] == "kk", str(pdf))
    docx = rag.ingest(owner, "Instr.docx", (SAMPLES / "Instrukciya_po_ohrane_truda.docx").read_bytes())
    check("DOCX + таблица", docx["chunks"] >= 2 and docx["lang"] == "ru", str(docx))
    txt = rag.ingest(owner, "note.txt", "Wi-Fi пароль в офисе: Talap2026. Обед с 13:00 до 14:00.".encode("cp1251"))
    check("TXT (cp1251)", txt["chars"] > 20)

    print(f"OCR ({ocr_engine()}):")
    if ocr_engine() != "none":
        scan = rag.ingest(owner, "scan.pdf", (SAMPLES / "Scan_AIChatBots_talaptary.pdf").read_bytes())
        check("скан распознан", scan["ocr_pages"] == 1 and scan["chars"] > 100, str(scan))
        scan_text = " ".join(c["text"] for c in store.doc_chunks({scan["id"]}))
        check("текст скана", "балл" in scan_text.lower(), scan_text[:120])
    else:
        print("  ⚠️  OCR недоступен — пропуск")

    print("Поиск и офлайн-ответы:")
    settings.full_context_chars = 0  # принудительно режим поиска (BM25)
    cases = [
        ("Финалдық кезең қашан өтеді?", "14 қараша"),
        ("Команда құрамында неше адам болуы мүмкін?", "үш адамнан"),
        ("Как часто проводится повторный инструктаж?", "6 месяцев"),
        ("Срок носки защитной каски?", "24 месяца"),
        ("Какой пароль от Wi-Fi?", "Talap2026"),
    ]
    for question, expected in cases:
        result = await ask(question, owner)
        ok = expected.lower() in result["text"].lower() and bool(result["sources"])
        loc = result["sources"][0]["loc_label"] if result["sources"] else "-"
        check(question, ok, f"источник: {result['sources'][0]['doc_name'] if result['sources'] else '-'} {loc}")

    print("Смысловой поиск (эмбеддинги):")
    import time
    from app.semantic import semantic
    semantic.start(lambda: list(store.chunks))
    for _ in range(600):
        if semantic.ready or semantic.error:
            break
        time.sleep(0.5)
    if semantic.ready:
        allowed = {pdf["id"], docx["id"]}
        for question, expected in [("Как часто заново учат технике безопасности?", "6 месяцев"),
                                   ("What should I do if a fire starts?", "пожар")]:
            hits = rag._search([question], allowed, 3)
            check(f"«{question}» → «{expected}»",
                  any(expected in c["text"] for c, _ in hits), str([c["text"][:50] for c, _ in hits]))
    else:
        print(f"  ⚠️  модель недоступна ({semantic.error}) — пропуск")

    print("Сравнение версий:")
    import io
    import docx as docx_lib
    from app import compare
    d = docx_lib.Document(io.BytesIO((SAMPLES / "Instrukciya_po_ohrane_truda.docx").read_bytes()))
    for p in d.paragraphs:
        if "6 месяцев" in p.text:
            for r in p.runs:
                r.text = r.text.replace("6 месяцев", "3 месяца")
    d.add_paragraph("6. Новый раздел: ежегодная аттестация обязательна.")
    buf = io.BytesIO()
    d.save(buf)
    docx2 = rag.ingest(owner, "Instr_v2.docx", buf.getvalue())
    diff = compare.diff_docs(docx, docx2, "ru")
    words = [w for c in diff["changes"] if c["kind"] == "changed" for w in c["words"] if w[0] != "eq"]
    check("изменение найдено по словам (6 → 3)", ["del", "6"] in words and ["ins", "3"] in words, str(words))
    check("добавленный пункт", any(c["kind"] == "added" and "аттестация" in c["new"] for c in diff["changes"]))
    check("остальное совпадает", diff["stats"]["similarity"] >= 85 and diff["stats"]["removed"] == 0, str(diff["stats"]))
    same = compare.diff_docs(docx, docx, "ru")
    check("одинаковые версии → нет изменений", same["changes"] == [])
    store.delete(docx2["id"])

    result = await ask("Финал қашан?", "nobody")
    check("нет документов → подсказка", "құжат" in result["text"].lower())

    print("REST API:")
    client = TestClient(app)
    headers = {}
    check("без входа — 401", client.get("/api/documents").status_code == 401)
    check("сравнение без входа — 401", client.get(f"/api/compare?a={pdf['id']}&b={docx['id']}").status_code == 401)
    check("без входа — главная видна, данные закрыты",
          client.get("/", follow_redirects=False).status_code == 200 and client.get("/api/chats").status_code == 401)
    from app import mailer
    from app.auth import normalize_email
    check("Gmail-алиасы = один адрес", normalize_email("Ai.Gerim+test@GMAIL.com") == "aigerim@gmail.com")
    reg = {"name": "Айгерім", "email": "aigerim@example.com", "password": "secret2026", "lang": "ru"}
    check("слабый пароль отклонён", client.post("/api/auth/register", json={**reg, "password": "onlyletters"}).status_code == 400)
    check("одноразовая почта отклонена",
          client.post("/api/auth/register", json={**reg, "email": "x@mailinator.com"}).status_code == 400)
    r = client.post("/api/auth/register", json=reg)
    check("регистрация → нужен код", r.status_code == 200 and r.json()["need_verify"] is True, r.text[:120])
    check("без подтверждения вход запрещён",
          client.post("/api/auth/login", json={"email": reg["email"], "password": reg["password"]}).status_code == 403)
    check("неверный код", client.post("/api/auth/verify", json={"email": reg["email"], "code": "000000"}).status_code == 400)
    code = [m["code"] for m in mailer.outbox if m["to"] == reg["email"] and m["kind"] == "verify"][-1]
    r = client.post("/api/auth/verify", json={"email": reg["email"], "code": code})
    check("подтверждение email кодом", r.status_code == 200 and client.get("/api/auth/me").status_code == 200)
    check("повторная регистрация того же email",
          client.post("/api/auth/register", json={**reg, "email": "AIGERIM@example.com"}).status_code == 400)
    stranger = TestClient(app)
    check("неверный пароль", stranger.post("/api/auth/login", json={"email": reg["email"], "password": "wrong1234"}).status_code == 401)
    check("вход по паролю", stranger.post("/api/auth/login", json={"email": reg["email"], "password": reg["password"]}).status_code == 200)
    stranger.post("/api/auth/logout")
    check("выход", stranger.get("/api/auth/me").status_code == 401)
    check("сброс: ответ одинаковый для неизвестного email",
          stranger.post("/api/auth/send-code", json={"email": "nobody@example.com", "kind": "reset"}).status_code == 200)
    stranger.post("/api/auth/send-code", json={"email": reg["email"], "kind": "reset"})
    code = [m["code"] for m in mailer.outbox if m["kind"] == "reset"][-1]
    r = stranger.post("/api/auth/reset", json={"email": reg["email"], "code": code, "password": "newpass2026"})
    check("сброс пароля по коду", r.status_code == 200)
    check("старая сессия закрыта после сброса", client.get("/api/auth/me").status_code == 401)
    client.post("/api/auth/login", json={"email": reg["email"], "password": "newpass2026"})
    check("вход с новым паролем", client.get("/api/auth/me").status_code == 200)
    r = client.put("/api/settings", json={"theme": "dark", "answer_style": "concise", "show_trust": False, "theme_bad": 1})
    check("настройки сохраняются", r.status_code == 200 and r.json()["theme"] == "dark" and r.json()["show_trust"] is False
          and "theme_bad" not in r.json())
    stranger.post("/api/auth/logout")
    from app.auth import users as user_store, AuthError
    fb_user = user_store.firebase_login("fb-uid-1", "Ali.Test@gmail.com", "Али")
    check("вход через Firebase создаёт аккаунт", fb_user["verified"] and fb_user["email_key"] == "alitest@gmail.com")
    check("повторный вход Firebase — тот же аккаунт", user_store.firebase_login("fb-uid-1", "ali.test@gmail.com", "")["id"] == fb_user["id"])
    try:
        user_store.firebase_login("fb-uid-2", "alitest+x@gmail.com", "")
        check("Gmail-алиас другого аккаунта отклонён", False)
    except AuthError:
        check("Gmail-алиас другого аккаунта отклонён", True)
    check("проверка email до регистрации",
          client.post("/api/auth/check-email", json={"email": "ali.test+1@gmail.com"}).status_code == 400
          and client.post("/api/auth/check-email", json={"email": "fresh@example.com"}).status_code == 200)
    check("гостевой вход отключён", stranger.post("/api/auth/guest").status_code in (404, 405))
    other = {"name": "Бауыржан", "email": "baur@example.com", "password": "other2026x"}
    stranger.post("/api/auth/register", json=other)
    code = [m["code"] for m in mailer.outbox if m["to"] == other["email"]][-1]
    check("второй пользователь", stranger.post("/api/auth/verify", json={"email": other["email"], "code": code}).status_code == 200)
    r = client.get("/api/status")
    check("/api/status", r.status_code == 200 and r.json()["llm_enabled"] is False)
    with open(SAMPLES / "Instrukciya_po_ohrane_truda.docx", "rb") as f:
        r = client.post("/api/documents", headers=headers, files={"files": ("instr.docx", f)})
    check("загрузка через API", r.status_code == 200 and "id" in r.json()[0], r.text[:200])
    doc_id = r.json()[0]["id"]
    r = client.post("/api/ask", headers=headers, json={"question": "Какой номер вызова пожарной службы?"})
    check("стриминг /api/ask", r.status_code == 200 and "101" in r.text, r.text[:150])
    chat_list = client.get("/api/chats").json()
    check("чат сохранён после вопроса", len(chat_list) == 1 and "пожар" in chat_list[0]["title"].lower(), str(chat_list)[:120])
    chat_id = chat_list[0]["id"]
    r = client.post("/api/ask", headers=headers, json={"question": "А скорая помощь?", "chat_id": chat_id})
    msgs = client.get(f"/api/chats/{chat_id}").json()["messages"]
    check("продолжение того же чата", len(msgs) == 4 and msgs[1]["role"] == "bot" and len(client.get("/api/chats").json()) == 1)
    check("чужой чат недоступен", TestClient(app).get(f"/api/chats/{chat_id}").status_code in (401, 404))
    check("удаление чата", client.delete(f"/api/chats/{chat_id}").status_code == 200 and client.get("/api/chats").json() == [])
    r = client.get(f"/api/documents/{doc_id}/summary?lang=ru", headers=headers)
    check("резюме", r.status_code == 200 and r.json()["summary"])
    r = client.get(f"/api/documents/{doc_id}/insights?lang=ru", headers=headers)
    check("карта документа", r.status_code == 200 and len(r.json()["numbers"]) >= 3 and r.json()["requirements"])
    r = client.get(f"/api/documents/{doc_id}/quiz?lang=ru&n=3", headers=headers)
    quiz = r.json().get("questions", [])
    check("тест по документу", r.status_code == 200 and len(quiz) >= 2
          and all(q["options"][q["answer"]] in q["quote"] for q in quiz), str(quiz[:1])[:150])
    r = client.get(f"/api/documents/{doc_id}/cards?lang=ru&n=5", headers=headers)
    cards = r.json().get("cards", [])
    check("карточки", r.status_code == 200 and len(cards) >= 2 and all(c["front"] and c["back"] for c in cards),
          str(cards[:1])[:150])
    for fmt, magic in (("pdf", b"%PDF"), ("docx", b"PK")):
        r = client.post("/api/export", headers=headers, json={
            "format": fmt, "kind": "answer", "lang": "kk", "question": "Каска мерзімі?",
            "answer": "**24 ай** [1]\n\n| Құрал | Мерзім |\n|---|---|\n| Каска | 24 ай |\n\n- Қазақша әріптер: әғқңөұүһі",
            "sources": [{"n": 1, "doc_name": "instr.docx", "loc_label": "бөлім 1", "text": "Каска — 24 месяца"}]})
        check(f"экспорт ответа в {fmt.upper()}", r.status_code == 200 and r.content[:4].startswith(magic)
              and "attachment" in r.headers.get("content-disposition", ""), r.text[:120] if r.status_code != 200 else "")
    r = stranger.post("/api/export", json={"format": "pdf", "kind": "compare", "a": doc_id, "b": doc_id})
    check("чужое сравнение не экспортируется", r.status_code == 404)
    with open(SAMPLES / "Innovation_AI_Hackathon_2026_Ereje.pdf", "rb") as f:
        pdf_id = client.post("/api/documents", headers=headers, files={"files": ("ereje.pdf", f)}).json()[0]["id"]
    chunk_id = f"{pdf_id}:0"
    r = client.get(f"/api/chunk?chunk={chunk_id}", headers=headers)
    check("фрагмент-доказательство", r.status_code == 200 and r.json()["has_page"])
    r = client.get(f"/api/evidence.png?chunk={chunk_id}&q=Жамбыл облысы")
    check("страница с подсветкой", r.status_code == 200 and r.content[:4] == b"\x89PNG")
    r = stranger.get(f"/api/evidence.png?chunk={chunk_id}")
    check("чужая страница недоступна", r.status_code == 404)
    client.delete(f"/api/documents/{pdf_id}", headers=headers)
    r = client.get(f"/api/documents/{doc_id}/file")
    check("скачивание оригинала", r.status_code == 200 and len(r.content) > 1000)
    r = stranger.get(f"/api/documents/{doc_id}/file")
    check("чужой документ недоступен", r.status_code == 404)
    r = client.post("/api/documents", headers=headers, files={"files": ("virus.exe", b"MZ")})
    check("неподдерживаемый формат", "error" in r.json()[0])
    r = client.delete(f"/api/documents/{doc_id}", headers=headers)
    check("удаление", r.status_code == 200 and client.get("/api/documents", headers=headers).json() == [])

    print("\nИТОГ:", "все проверки пройдены ✅" if not failures else f"ошибок: {failures} ❌")
    sys.exit(1 if failures else 0)


asyncio.run(main())
