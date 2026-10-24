"""Экспорт ответа или сравнения версий в PDF и DOCX — чтобы отправить коллеге или руководителю.

Содержимое описывается списком блоков (заголовок, абзац, пункт списка, таблица), абзац — списком
кусочков текста со стилем. Из одного описания собираются оба формата:
  * DOCX — python-docx;
  * PDF  — PyMuPDF Story (HTML → страницы A4) со шрифтом, в котором есть казахские буквы.
"""
from __future__ import annotations

import html
import io
import re
import time
from pathlib import Path

LABELS = {
    "answer_t": {"kk": "Mind AI жауабы", "ru": "Ответ Mind AI", "en": "Mind AI answer"},
    "compare_t": {"kk": "Нұсқаларды салыстыру", "ru": "Сравнение версий", "en": "Version comparison"},
    "question": {"kk": "Сұрақ", "ru": "Вопрос", "en": "Question"},
    "answer": {"kk": "Жауап", "ru": "Ответ", "en": "Answer"},
    "sources": {"kk": "Дереккөздер", "ru": "Источники", "en": "Sources"},
    "points": {"kk": "Негізгі өзгерістер", "ru": "Главные изменения", "en": "Key changes"},
    "attention": {"kk": "Назар аударыңыз", "ru": "Обратите внимание", "en": "Pay attention"},
    "all": {"kk": "Барлық өзгерістер", "ru": "Все изменения", "en": "All changes"},
    "added": {"kk": "Қосылды", "ru": "Добавлено", "en": "Added"},
    "removed": {"kk": "Өшірілді", "ru": "Удалено", "en": "Removed"},
    "changed": {"kk": "Өзгерді", "ru": "Изменено", "en": "Changed"},
    "similar": {"kk": "сәйкестік", "ru": "совпадение", "en": "match"},
    "footer": {"kk": "Mind AI — құжаттар бойынша интеллектуалды көмекші",
               "ru": "Mind AI — интеллектуальный помощник по документам",
               "en": "Mind AI — intelligent document assistant"},
}

# Шрифты с кириллицей и казахскими буквами: в Docker — DejaVu (fonts-dejavu-core), на macOS — Arial
FONTS = [
    ("/usr/share/fonts/truetype/dejavu", "DejaVuSans.ttf", "DejaVuSans-Bold.ttf"),
    ("/System/Library/Fonts/Supplemental", "Arial.ttf", "Arial Bold.ttf"),
]


def _l(key: str, lang: str) -> str:
    return LABELS[key].get(lang, LABELS[key]["ru"])


# ---------------------------------------------------------------- markdown → блоки

def _inline(text: str) -> list[tuple[str, str]]:
    """**жирный**, `код` и ссылки на источники [1] → [(текст, стиль)], стиль: "" | "b" | "sup"."""
    runs = []
    for part in re.split(r"(\*\*[^*]+\*\*|\[\d+\]|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            runs.append((part[2:-2], "b"))
        elif re.fullmatch(r"\[\d+\]", part):
            runs.append((part, "sup"))
        elif part.startswith("`"):
            runs.append((part[1:-1], ""))
        else:
            runs.append((part.replace("*", "") if part.count("*") % 2 == 0 else part, ""))
    return runs


def md_blocks(text: str) -> list[tuple]:
    blocks, para, table = [], [], []

    def flush():
        if para:
            blocks.append(("p", _inline(" ".join(para))))
            para.clear()
        if table:
            blocks.append(("table", [list(r) for r in table]))
            table.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if line.startswith("|"):
            if para:
                blocks.append(("p", _inline(" ".join(para))))
                para.clear()
            cells = [c.strip() for c in line.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                table.append(cells)
            continue
        if not line:
            flush()
            continue
        if table:
            flush()
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            flush()
            blocks.append(("h3", _inline(m.group(2))))
        elif re.match(r"^[-*•]\s+", line):
            flush()
            blocks.append(("li", _inline(re.sub(r"^[-*•]\s+", "", line))))
        elif re.match(r"^\d+[.)]\s+", line):
            flush()
            blocks.append(("li", _inline(line)))
        else:
            para.append(line)
    flush()
    return blocks


# ---------------------------------------------------------------- содержимое

def answer_blocks(question: str, answer: str, sources: list[dict], lang: str) -> tuple[str, list[tuple]]:
    blocks = [("h2", [(_l("question", lang), "")]), ("p", [(question, "")]),
              ("h2", [(_l("answer", lang), "")]), *md_blocks(answer)]
    if sources:
        blocks.append(("h2", [(_l("sources", lang), "")]))
        for s in sources:
            snippet = re.sub(r"\s+", " ", s.get("text", "")).strip()
            snippet = snippet[:220] + "…" if len(snippet) > 220 else snippet
            where = " · ".join(x for x in (s.get("doc_name", ""), s.get("loc_label", "")) if x)
            runs = [(f"[{s.get('n', '')}] ", "b"), (where, "b")]
            if snippet:
                runs.append((f" — «{snippet}»", "muted"))
            blocks.append(("li", runs))
    return _l("answer_t", lang), blocks


def compare_blocks(result: dict, lang: str) -> tuple[str, list[tuple]]:
    st, sm = result["stats"], result.get("summary")
    blocks = [("p", [(result["a"]["name"], "b"), ("  →  ", ""), (result["b"]["name"], "b")])]
    if sm:
        blocks.append(("p", [(sm["verdict"], "")]))
    blocks.append(("p", [(f"+{st['added']} {_l('added', lang).lower()}   −{st['removed']} {_l('removed', lang).lower()}   "
                          f"{st['changed']} {_l('changed', lang).lower()}   {st['similarity']}% {_l('similar', lang)}", "muted")]))
    for key in ("points", "attention"):
        if sm and sm.get(key):
            blocks.append(("h2", [(_l(key, lang), "")]))
            blocks += [("li", [(p, "")]) for p in sm[key]]
    blocks.append(("h2", [(_l("all", lang), "")]))
    for c in result["changes"]:
        where = c.get("loc_new") or c.get("loc_old") or ""
        blocks.append(("tag", [(_l(c["kind"], lang).upper(), c["kind"]), (f"  {where}", "muted")]))
        if c["kind"] == "changed":
            runs = [(t, {"eq": "", "del": "del", "ins": "ins"}[op]) for op, t in c["words"]]
        elif c["kind"] == "added":
            runs = [(c["new"], "ins")]
        else:
            runs = [(c["old"], "del")]
        blocks.append(("p", runs))
    return _l("compare_t", lang), blocks


# ---------------------------------------------------------------- DOCX

def to_docx(title: str, blocks: list[tuple], lang: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    doc.add_heading(title, level=1)
    meta = doc.add_paragraph()
    run = meta.add_run(f"{_l('footer', lang)} · {time.strftime('%d.%m.%Y %H:%M')}")
    run.font.size, run.font.color.rgb = Pt(9), RGBColor(0x80, 0x80, 0x88)

    def fill(par, runs):
        for text, style in runs:
            r = par.add_run(text)
            if style in ("b", "added", "removed", "changed"):
                r.bold = True
            if style == "sup":
                r.font.superscript = True
                r.font.color.rgb = RGBColor(0x4F, 0x52, 0xD1)
            elif style == "muted":
                r.font.color.rgb = RGBColor(0x70, 0x70, 0x78)
            elif style in ("del", "removed"):
                r.font.color.rgb = RGBColor(0xC0, 0x30, 0x3C)
                r.font.strike = style == "del"
            elif style in ("ins", "added"):
                r.font.color.rgb = RGBColor(0x14, 0x7A, 0x48)
                r.underline = style == "ins"
            elif style == "changed":
                r.font.color.rgb = RGBColor(0xA8, 0x67, 0x0C)

    for kind, content in blocks:
        if kind == "h2":
            doc.add_heading(content[0][0], level=2)
        elif kind == "h3":
            fill(doc.add_heading("", level=3), content)
        elif kind == "li":
            fill(doc.add_paragraph(style="List Bullet"), content)
        elif kind == "table":
            width = max(len(r) for r in content)
            table = doc.add_table(rows=0, cols=width)
            table.style = "Table Grid"
            for i, row in enumerate(content):
                cells = table.add_row().cells
                for j in range(width):
                    par = cells[j].paragraphs[0]
                    fill(par, [(t, "b" if i == 0 else s) for t, s in _inline(row[j] if j < len(row) else "")])
        elif kind == "tag":
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            fill(par, content)
        else:
            fill(doc.add_paragraph(), content)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

CSS = """
@font-face {{ font-family: Body; src: url("{regular}"); }}
@font-face {{ font-family: Body; src: url("{bold}"); font-weight: bold; }}
* {{ font-family: Body; }}
body {{ font-size: 10.5pt; line-height: 1.45; color: #1c1c22; }}
h1 {{ font-size: 19pt; margin: 0 0 2pt 0; }}
h2 {{ font-size: 13pt; margin: 14pt 0 4pt 0; color: #33337a; }}
h3 {{ font-size: 11.5pt; margin: 10pt 0 3pt 0; }}
p {{ margin: 0 0 6pt 0; }}
ul {{ margin: 0 0 6pt 0; }}
li {{ margin: 0 0 3pt 0; }}
table {{ border-collapse: collapse; margin: 4pt 0 8pt 0; }}
td {{ border: 0.6pt solid #b8b8c4; padding: 3pt 5pt; vertical-align: top; }}
.meta, .muted {{ color: #737380; }}
.meta {{ font-size: 8.5pt; margin-bottom: 10pt; }}
.sup {{ color: #4f52d1; font-size: 7.5pt; }}
.del {{ color: #c0303c; text-decoration: line-through; }}
.ins {{ color: #147a48; text-decoration: underline; }}
.tag {{ margin: 8pt 0 1pt 0; font-size: 8.5pt; }}
.added {{ color: #147a48; font-weight: bold; }}
.removed {{ color: #c0303c; font-weight: bold; }}
.changed {{ color: #a8670c; font-weight: bold; }}
"""


def _font_dir() -> tuple[str, str, str] | None:
    for folder, regular, bold in FONTS:
        if (Path(folder) / regular).exists() and (Path(folder) / bold).exists():
            return folder, regular, bold
    return None


def _html_runs(runs: list[tuple[str, str]]) -> str:
    out = []
    for text, style in runs:
        t = html.escape(text)
        if style == "b":
            out.append(f"<b>{t}</b>")
        elif style:
            out.append(f'<span class="{style}">{t}</span>')
        else:
            out.append(t)
    return "".join(out)


def to_pdf(title: str, blocks: list[tuple], lang: str) -> bytes:
    import pymupdf

    fonts = _font_dir()
    if not fonts:
        raise RuntimeError("Не найден шрифт с кириллицей для PDF (установите fonts-dejavu-core)")
    parts = [f"<h1>{html.escape(title)}</h1>",
             f'<p class="meta">{html.escape(_l("footer", lang))} · {time.strftime("%d.%m.%Y %H:%M")}</p>']
    in_list = False
    for kind, content in blocks:
        if kind == "li" and not in_list:
            parts.append("<ul>")
            in_list = True
        elif kind != "li" and in_list:
            parts.append("</ul>")
            in_list = False
        if kind in ("h2", "h3"):
            parts.append(f"<{kind}>{_html_runs(content)}</{kind}>")
        elif kind == "li":
            parts.append(f"<li>{_html_runs(content)}</li>")
        elif kind == "table":
            rows = []
            for i, row in enumerate(content):
                cells = "".join(f"<td>{'<b>' if i == 0 else ''}{_html_runs(_inline(c))}{'</b>' if i == 0 else ''}</td>"
                                for c in row)
                rows.append(f"<tr>{cells}</tr>")
            parts.append(f"<table>{''.join(rows)}</table>")
        elif kind == "tag":
            parts.append(f'<p class="tag">{_html_runs(content)}</p>')
        else:
            parts.append(f"<p>{_html_runs(content)}</p>")
    if in_list:
        parts.append("</ul>")

    folder, regular, bold = fonts
    story = pymupdf.Story(html="".join(parts), user_css=CSS.format(regular=regular, bold=bold),
                          archive=pymupdf.Archive(folder))
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    page = pymupdf.paper_rect("a4")
    area = page + (54, 54, -54, -54)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()
    return buf.getvalue()


def build(fmt: str, title: str, blocks: list[tuple], lang: str) -> bytes:
    return to_pdf(title, blocks, lang) if fmt == "pdf" else to_docx(title, blocks, lang)
