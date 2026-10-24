"""Извлечение текста из документов разных форматов.

Каждый парсер возвращает список «страниц»: {"loc": {"type": ..., "value": ...}, "text": str, "ocr": bool}
loc нужен, чтобы в ответе показывать точный источник: страница PDF, слайд, лист Excel и т.д.
"""
import csv
import io
import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

from .ocr import ocr_available, ocr_bytes, ocr_image

log = logging.getLogger(__name__)

TEXT_EXT = {".txt", ".md", ".markdown", ".json", ".xml", ".log", ".rst"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
TEXTUTIL_EXT = {".doc", ".rtf", ".odt"}
SUPPORTED_EXT = {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".html", ".htm"} | TEXT_EXT | IMAGE_EXT | TEXTUTIL_EXT

# Страница PDF, на которой меньше символов — считаем сканом и отправляем в OCR
SCAN_THRESHOLD = 40


class UnsupportedFormat(ValueError):
    pass


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1251", "latin-1"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "ignore")


def _page(kind: str, value, text: str, ocr: bool = False) -> dict:
    return {"loc": {"type": kind, "value": value}, "text": text.strip(), "ocr": ocr}


def parse_pdf(data: bytes) -> list[dict]:
    import pymupdf

    pages = []
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text") or ""
            used_ocr = False
            if len(text.strip()) < SCAN_THRESHOLD and ocr_available():
                # Скан: рендерим страницу в картинку 300 dpi и распознаём
                pix = page.get_pixmap(dpi=300)
                try:
                    ocr_text = ocr_bytes(pix.tobytes("png"))
                except Exception as exc:  # noqa: BLE001
                    log.warning("OCR страницы %s не удался: %s", i, exc)
                    ocr_text = ""
                if len(ocr_text.strip()) > len(text.strip()):
                    text, used_ocr = ocr_text, True
            pages.append(_page("page", i, text, used_ocr))
    return pages


def parse_docx(data: bytes) -> list[dict]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    # Обходим тело документа по порядку, чтобы таблицы оказались на своих местах
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            if style.startswith("heading") or style.startswith("title") or style.startswith("заголовок"):
                lines.append(f"\n## {text}")
            elif "list" in style:
                lines.append(f"• {text}")
            else:
                lines.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    value = cell.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != value:  # объединённые ячейки повторяются
                        cells.append(value)
                if any(cells):
                    lines.append(" | ".join(cells))
    # Колонтитулы тоже бывают полезны (название организации, номер приказа)
    for section in document.sections:
        for para in section.header.paragraphs:
            if para.text.strip():
                lines.insert(0, para.text.strip())
                break

    # Картинки внутри DOCX (например, вставленный скан) — распознаём
    pages = _split_sections("\n".join(lines))
    if ocr_available():
        n = 0
        for rel in document.part.rels.values():
            if "image" in rel.reltype:
                try:
                    text = ocr_bytes(rel.target_part.blob)
                except Exception:  # noqa: BLE001
                    continue
                if len(text.strip()) >= SCAN_THRESHOLD:
                    n += 1
                    pages.append(_page("image", n, text, True))
    return pages


def _split_sections(text: str, size: int = 3000) -> list[dict]:
    """Документы без страниц (DOCX, TXT) делим на условные разделы ~3000 символов по абзацам."""
    paragraphs = [p for p in text.split("\n") if p.strip()]
    sections, buf, length = [], [], 0
    for para in paragraphs:
        if length + len(para) > size and buf:
            sections.append("\n".join(buf))
            buf, length = [], 0
        buf.append(para)
        length += len(para) + 1
    if buf:
        sections.append("\n".join(buf))
    return [_page("section", i, s) for i, s in enumerate(sections, start=1)]


def parse_pptx(data: bytes) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Заметки: " + slide.notes_slide.notes_text_frame.text.strip())
        pages.append(_page("slide", i, "\n".join(parts)))
    return pages


def parse_xlsx(data: bytes) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            values = ["" if v is None else str(v).strip() for v in row]
            if any(values):
                rows.append(" | ".join(values).rstrip(" |"))
        pages.append(_page("sheet", ws.title, "\n".join(rows)))
    wb.close()
    return pages


def parse_csv(data: bytes) -> list[dict]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = [" | ".join(r) for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    return [_page("sheet", 1, "\n".join(rows))]


def parse_html(data: bytes) -> list[dict]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
    return _split_sections(text)


def parse_textutil(data: bytes, ext: str) -> list[dict]:
    """DOC / RTF / ODT — конвертируем встроенной утилитой macOS textutil или LibreOffice."""
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / f"input{ext}"
        src.write_bytes(data)
        if shutil.which("textutil"):
            out = Path(tmp) / "input.txt"
            subprocess.run(["textutil", "-convert", "txt", "-output", str(out), str(src)],
                           check=True, capture_output=True, timeout=60)
            return _split_sections(out.read_text("utf-8", "ignore"))
        office = shutil.which("soffice") or shutil.which("libreoffice")
        if office:
            subprocess.run([office, "--headless", "--convert-to", "txt:Text", "--outdir", tmp, str(src)],
                           check=True, capture_output=True, timeout=120)
            return _split_sections((Path(tmp) / "input.txt").read_text("utf-8", "ignore"))
    raise UnsupportedFormat(f"Для формата {ext} нужен LibreOffice. Сохраните файл как DOCX или PDF.")


def parse_image(data: bytes) -> list[dict]:
    if not ocr_available():
        raise UnsupportedFormat("OCR недоступен: установите Tesseract или включите VISION_OCR")
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        text = ocr_image(img)
    return [_page("image", 1, text, True)]


def parse_file(filename: str, data: bytes) -> list[dict]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXT:
        raise UnsupportedFormat(f"Формат {ext or '(без расширения)'} не поддерживается")
    if ext == ".pdf":
        pages = parse_pdf(data)
    elif ext == ".docx":
        pages = parse_docx(data)
    elif ext == ".pptx":
        pages = parse_pptx(data)
    elif ext == ".xlsx":
        pages = parse_xlsx(data)
    elif ext == ".csv":
        pages = parse_csv(data)
    elif ext in {".html", ".htm"}:
        pages = parse_html(data)
    elif ext in IMAGE_EXT:
        pages = parse_image(data)
    elif ext in TEXTUTIL_EXT:
        pages = parse_textutil(data, ext)
    else:
        pages = _split_sections(_decode(data))
    return [p for p in pages if p["text"]]
